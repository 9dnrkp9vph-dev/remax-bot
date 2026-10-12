# tests/deals_history/test_import_tool.py
import sys, os, datetime as dt
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))
import openpyxl
from tools.import_deal_history import rows_from_workbook
n = 0
def ck(c, m):
    global n; assert c, m; n += 1
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "עסקאות25"
ws.append(["דוח"]); ws.append(['מס"ד', "תאריך עסקה", "עיר", "רחוב", "קונה / מוכר", "סוכן", "סכום העסקה", "עמלה לפני מעמ", "התחשבנות סוכן"])
ws.append([1, dt.datetime(2025, 1, 2), "קרית ים", "הרצל 1", "מוכר", "משה", 1000000, 10000, 5000])
ws.append([2, dt.datetime(2025, 1, 3), "קרית ים", "הרצל 1", "קונה", "משה", 1000000, 9000, None])
ws.append([None, None, None, None, None, 'סה"כ', None, 19000, None])
ws.append([3, None, None, None, None, "דנה", None, None, None])
lawyer = wb.create_sheet('דוח לעו"ד')
lawyer.append(["תאריך", "כתובת", "מתווך", "סכום עסקה", "התחשבנות", "עו\"ד"]); lawyer.append([dt.datetime(2025, 1, 2), "הרצל 1", "משה", 1000000, 1, "x"])
ok, rep = rows_from_workbook(wb, "kiryat_yam", 2025, ["עסקאות25"])
ck(len(ok) == 2 and rep["ok"] == 2 and rep["skip"] >= 1 and rep["broken"] == 1, rep)
ck(ok[0]["src_key"] == "kiryat_yam|2025|עסקאות25|1" and ok[0]["raw"].get("התחשבנות סוכן") == "5000", ok[0])
ck(rep["agents"] == {"משה": 2}, rep["agents"])
ck(all(r["src_sheet"] == "עסקאות25" for r in ok), "only approved sheet (no דוח לעו\"ד)")
ok2, rep2 = rows_from_workbook(wb, "kiryat_yam", 2025, ["אין כזה"])
ck(ok2 == [] and rep2["missing_sheets"] == ["אין כזה"], rep2)
ws2 = openpyxl.Workbook().active; ws2.title = "מעקב עסקאות 2023"
ws2.append(["תאריך עסקה", "עיר", "רחוב", "קונה / מוכר", "סוכן", "סכום העסקה"])
ws2.append([dt.datetime(2025, 2, 2), "קרית ים", "הגפן 1", "מוכר", None, 800000])
ok3, rep3 = rows_from_workbook(ws2.parent, "kiryat_yam", 2025, ["מעקב עסקאות 2023"], default_agent="קובי חמו")
ck(len(ok3) == 1 and ok3[0]["agent"] == "קובי חמו", ("default agent", rep3))
print("OK", n, "checks")
