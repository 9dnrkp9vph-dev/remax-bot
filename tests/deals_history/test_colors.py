# tests/deals_history/test_colors.py — אייל 12/10: 'המון טרם שולם למרות שבאקסל רואים שכסף התקבל'
# בקובץ ביאליק 2026 'התקבל כסף'/'הופקד'/'דווח במערכת' מסומנים בצבע התא (ירוק C6EFCE / ורוד FFC7CE / צהוב FFEB9C)
import sys, os, io, datetime as dt
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))
import openpyxl
from openpyxl.styles import PatternFill
import deals_history as D
n = 0
def ck(c, m):
    global n; assert c, m; n += 1
G, P, Y = (PatternFill("solid", fgColor=c) for c in ("C6EFCE", "FFC7CE", "FFEB9C"))
ck(D.fill_class_rgb("FFC6EFCE") == "green" and D.fill_class_rgb("C6EFCE") == "green" and D.fill_class_rgb("FFFFC7CE") == "red"
   and D.fill_class_rgb("FFFFEB9C") == "yellow" and D.fill_class_rgb("00000000") == "" and D.fill_class_rgb(None) == "", "classes")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "עסקאות 26"
ws.append(['מס"ד', "תאריך עסקה", "עיר", "רחוב", "קונה / מוכר", "סוכן", "סכום העסקה", "עמלה כולל מעמ", "עמלה לפני מעמ",
           "התקבל כסף", "הופקד", "דווח במערכת", "שם לקוח", "נייד לקוח"])
ws.append([13, dt.datetime(2026, 1, 5), "ביאליק", "שושנה דמרי 3", "מוכר", "חי", 2330000, 46600, 39491, None, None, dt.datetime(2026, 6, 1), "דמיטרי", 546334989])
ws.append([14, dt.datetime(2026, 2, 5), "ביאליק", "הרצל 2", "קונה", "חי", 1000000, 20000, 17000, None, None, None, "גל", None])
ws.append([15, dt.datetime(2026, 3, 5), "ביאליק", "הרצל 3", "קונה", "חי", 1000000, 20000, 17000, 5000, None, None, "רון", None])
ws.append([16, dt.datetime(2026, 4, 5), "ביאליק", "הרצל 4", "משכיר", "חי", 3500, None, 0, "לא משלם", None, None, "", None])
for col in "JKL":
    ws[col + "2"].fill = G
ws["J3"].fill = P; ws["K3"].fill = P; ws["L3"].fill = G       # לא שולם, אבל דווח (ירוק בלי תאריך)
ws["J4"].fill = P                                              # מספר בתא גובר על הצבע
ws["J5"].fill = Y
b = io.BytesIO(); wb.save(b)
ok, rep = D.rows_from_workbook(D.open_workbook(b.getvalue(), ""), "bialik", 2026, ["עסקאות 26"], key_mode="sync")
R = {r["src_serial"]: r for r in ok}
ck(R["13"]["paid_amount"] == 46600 and R["13"]["deposited"] == "כן" and R["13"]["reported_date"] == "2026-06-01", R["13"])
ck(R["13"]["client_phone"] == "0546334989", R["13"]["client_phone"])
ck(R["14"]["paid_amount"] is None and R["14"]["deposited"] == "" and R["14"]["reported_date"] is None
   and R["14"]["raw"]["_status"] == {"paid_amount": "red", "deposited": "red", "reported_date": "green"}, R["14"])
ck(R["15"]["paid_amount"] == 5000, "number wins")
ck(R["16"]["paid_amount"] is None, "yellow text = not paid")
# שינוי צבע בלבד משנה את ה-hash (אחרת הסנכרון לא מעדכן)
a = dict(R["14"]); b2 = dict(R["14"], raw=dict(R["14"]["raw"], _status={"paid_amount": "green"}))
ck(D.row_hash(a) != D.row_hash(b2), "color in hash")
# קריאה: 'דווח' ירוק בלי תאריך = דווח; לא בתור 'לא דווח'
rows = [dict(R["14"], id="14", status=R["14"]["raw"]["_status"]), dict(R["16"], id="16", status=R["16"]["raw"].get("_status"))]
ck([r["id"] for r in D.unreported(rows)] == ["16"], [r["id"] for r in D.unreported(rows)])
ck(D.view_agent(rows[0])["reported"] == "כן" and D.view_agent(rows[1])["reported"] == "", D.view_agent(rows[0]))
print("OK", n, "checks")
